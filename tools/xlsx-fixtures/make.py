"""Writes the .xlsx test files in app/src/sharedTest/fixtures/xlsx with real
spreadsheet writers, and the hostile files made from them.

    python3 tools/xlsx-fixtures/make.py

Needs LibreOffice (soffice), XlsxWriter and openpyxl. The README next to the
files records which program and version wrote each one. Running it again
rewrites every file; LibreOffice and XlsxWriter stamp the time into
docProps/core.xml, so the bytes change while the cells do not.
"""

import csv
import datetime
import io
import os
import pathlib
import shutil
import subprocess
import sys
import tempfile
import zipfile

import openpyxl
import xlsxwriter

ROOT = pathlib.Path(__file__).resolve().parents[2]
DESKTOP = ROOT / "app/src/sharedTest/fixtures/desktop"
OUT = ROOT / "app/src/sharedTest/fixtures/xlsx"


def libreoffice(csv_name: str, xlsx_name: str) -> None:
    """Converts one of the desktop app's own CSV fixtures with LibreOffice."""
    with tempfile.TemporaryDirectory() as work:
        source = pathlib.Path(work) / (pathlib.Path(xlsx_name).stem + ".csv")
        shutil.copy(DESKTOP / csv_name, source)
        subprocess.run(
            [
                "soffice", "--headless", "--norestore",
                "-env:UserInstallation=file://" + work + "/profile",
                # Comma separated, double quotes, UTF-8, from the first line.
                "--infilter=CSV:44,34,76,1",
                "--convert-to", "xlsx", "--outdir", work, str(source),
            ],
            check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )
        shutil.copy(pathlib.Path(work) / xlsx_name, OUT / xlsx_name)


def features() -> None:
    """Every kind of cell the reader turns into text, written by XlsxWriter."""
    book = xlsxwriter.Workbook(str(OUT / "features.xlsx"))
    date = book.add_format({"num_format": "yyyy-mm-dd"})
    stamp = book.add_format({"num_format": "yyyy-mm-dd hh:mm:ss"})
    clock = book.add_format({"num_format": "hh:mm"})
    custom = book.add_format({"num_format": 'd "of" mmmm yyyy'})
    money = book.add_format({"num_format": "#,##0.00 [$EUR]"})
    percent = book.add_format({"num_format": "0.0%"})
    bold = book.add_format({"bold": True})

    sheet = book.add_worksheet("Readings")
    sheet.write_row(0, 0, ["Station", "Reading", "Taken on", "Taken at", "Time", "Checked", "Ratio", "Cost", "Note"])
    sheet.write_string(1, 0, "Harbor Point")
    sheet.write_number(1, 1, 12.5)
    sheet.write_datetime(1, 2, datetime.date(2026, 1, 1), date)
    sheet.write_datetime(1, 3, datetime.datetime(2026, 1, 1, 6, 30, 15), stamp)
    sheet.write_datetime(1, 4, datetime.time(6, 30), clock)
    sheet.write_boolean(1, 5, True)
    sheet.write_number(1, 6, 0.125, percent)
    sheet.write_number(1, 7, 1234.5, money)
    sheet.write_string(1, 8, "First line\nsecond line")
    sheet.write_string(2, 0, "Ridge Field")
    sheet.write_number(2, 1, -3)
    sheet.write_datetime(2, 2, datetime.date(2024, 2, 29), custom)
    sheet.write_boolean(2, 5, False)
    sheet.write_formula(2, 6, "=B3/0", None, "#DIV/0!")
    sheet.write_formula(2, 7, "=B2*2", None, 25)
    sheet.write_formula(2, 8, '=A2&" north"', None, "Harbor Point north")
    # A gap of two rows and a cell far to the right.
    sheet.write_string(5, 0, "Lake Shore")
    sheet.write_number(5, 11, 1e-7)
    sheet.write_string(6, 0, 'Quotes "inside", and a comma')
    sheet.write_formula(6, 1, "=SUM(B2:B3)")  # XlsxWriter stores 0 as the result

    lookup = book.add_worksheet("Lookup")
    lookup.write_row(0, 0, ["Code", "Name"])
    lookup.write_row(1, 0, ["ST001", "Harbor Point"])
    lookup.hide()

    notes = book.add_worksheet("Notes")
    notes.write_rich_string(0, 0, "Plain, ", bold, "bold", " and plain again")
    notes.write_string(1, 0, "  leading and trailing spaces  ")
    notes.write_string(2, 0, "Tab\there")
    notes.write_string(3, 0, "Unicode: Zürich, 東京, مرحبا, 🙂")
    book.close()


def inline_strings() -> None:
    """XlsxWriter's constant-memory mode stores strings inline, not shared."""
    book = xlsxwriter.Workbook(str(OUT / "inline-strings.xlsx"), {"constant_memory": True})
    sheet = book.add_worksheet("Inline")
    with open(DESKTOP / "document-smoke.csv", newline="", encoding="utf-8") as source:
        for row, record in enumerate(csv.reader(source)):
            for column, value in enumerate(record):
                sheet.write_string(row, column, value)
    book.close()


def date_1904() -> None:
    book = xlsxwriter.Workbook(str(OUT / "date-1904.xlsx"), {"date_1904": True})
    sheet = book.add_worksheet("Dates")
    date = book.add_format({"num_format": "yyyy-mm-dd"})
    sheet.write_string(0, 0, "Day")
    sheet.write_datetime(1, 0, datetime.date(2026, 1, 1), date)
    sheet.write_datetime(2, 0, datetime.date(1904, 1, 2), date)
    book.close()


def by_openpyxl() -> None:
    """A second writer: openpyxl stores strings as shared strings."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Stations"
    with open(DESKTOP / "weather-stations.csv", newline="", encoding="utf-8") as source:
        for index, record in enumerate(csv.reader(source)):
            if index > 24:
                break
            sheet.append(record)
    sheet["J1"] = "Installed"
    sheet["J2"] = datetime.date(2019, 5, 17)
    sheet["J2"].number_format = "yyyy-mm-dd"
    sheet["K1"] = "Active"
    sheet["K2"] = True
    book.save(OUT / "openpyxl-stations.xlsx")


def large() -> None:
    """100,000 data rows: the desktop weather readings, repeated, in constant-memory mode."""
    with open(DESKTOP / "weather-stations.csv", newline="", encoding="utf-8") as source:
        records = list(csv.reader(source))
    header, rows = records[0], records[1:]
    book = xlsxwriter.Workbook(str(OUT / "large-weather.xlsx"), {"constant_memory": True})
    sheet = book.add_worksheet("Readings")
    sheet.write_row(0, 0, header)
    for index in range(100_000):
        sheet.write_row(index + 1, 0, rows[index % len(rows)])
    book.close()


def hostile() -> None:
    hostile_dir = OUT / "hostile"
    hostile_dir.mkdir(exist_ok=True)
    source = (OUT / "features.xlsx").read_bytes()

    # Cut in half, as an interrupted download leaves it.
    (hostile_dir / "truncated.xlsx").write_bytes(source[: len(source) // 2])

    # A DOCTYPE with an internal entity added to the shared strings.
    with zipfile.ZipFile(io.BytesIO(source)) as original, \
            zipfile.ZipFile(hostile_dir / "doctype.xlsx", "w", zipfile.ZIP_DEFLATED) as copy:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == "xl/sharedStrings.xml":
                data = data.replace(
                    b"?>",
                    b'?><!DOCTYPE sst [<!ENTITY a "aaaaaaaaaaaaaaaa">]>',
                    1,
                )
            copy.writestr(item, data)

    # The first worksheet replaced by 1.5 GiB of rows that deflate to about
    # 1.5 MB. The header of the entry tells the truth here; the reader must
    # stop by counting, not by trusting it.
    chunk = b"<row/>" * (1 << 20)
    total = 1536 * (1 << 20)
    with zipfile.ZipFile(io.BytesIO(source)) as original, \
            zipfile.ZipFile(hostile_dir / "zip-bomb.xlsx", "w", zipfile.ZIP_DEFLATED) as copy:
        for item in original.infolist():
            if item.filename == "xl/worksheets/sheet1.xml":
                entry = zipfile.ZipInfo(item.filename)
                entry.compress_type = zipfile.ZIP_DEFLATED
                with copy.open(entry, "w", force_zip64=True) as sheet:
                    sheet.write(b'<?xml version="1.0" encoding="UTF-8"?><worksheet><sheetData>')
                    written = 0
                    while written < total:
                        sheet.write(chunk)
                        written += len(chunk)
                    sheet.write(b"</sheetData></worksheet>")
            else:
                copy.writestr(item, original.read(item.filename))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    libreoffice("weather-stations.csv", "weather-stations.xlsx")
    libreoffice("document-smoke.csv", "document-smoke.xlsx")
    libreoffice("inline-quoted.csv", "inline-quoted.xlsx")
    features()
    inline_strings()
    date_1904()
    by_openpyxl()
    large()
    hostile()
    print("wrote", sorted(str(p.relative_to(OUT)) for p in OUT.rglob("*.xlsx")))


if __name__ == "__main__":
    sys.exit(main())
